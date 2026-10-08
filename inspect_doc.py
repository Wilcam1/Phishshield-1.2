import docx

def inspect_docx(filename):
    print(f"=== INSPECTING: {filename} ===")
    doc = docx.Document(filename)
    
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    print(f"Total sections: {len(doc.sections)}")
    
    headings = []
    sample_paragraphs = []
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style_name = p.style.name if p.style else "No style"
        if style_name.startswith("Heading") or style_name.startswith("Título") or len(text) < 80 and (text.isupper() or any(text.startswith(f"{n}.") for n in range(1, 10))):
            headings.append((i, style_name, text))
        if text:
            sample_paragraphs.append((i, style_name, text[:100]))
            
    print("\n--- DETECTED HEADINGS / KEY LINES ---")
    for idx, st, text in headings[:50]:
        print(f"[{idx}] ({st}): {text}")
    if len(headings) > 50:
        print(f"... and {len(headings) - 50} more headings.")

    print("\n--- FIRST 25 NON-EMPTY PARAGRAPHS ---")
    for idx, st, text in sample_paragraphs[:25]:
        print(f"[{idx}] ({st}): {text}")

if __name__ == "__main__":
    inspect_docx("Plantilla trabajo de grado 12.docx")

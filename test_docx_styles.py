import docx

doc = docx.Document("Plantilla trabajo de grado 12 - v2.docx")
print("Successfully opened Plantilla trabajo de grado 12 - v2.docx")
print(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

# Check style names
styles = set(p.style.name for p in doc.paragraphs if p.style)
print("Styles used in paragraphs:", sorted(list(styles)))

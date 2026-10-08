import docx

def dump_all_tables_info():
    doc = docx.Document("Plantilla trabajo de grado 12.docx")
    with open("tables_info_plantilla12.txt", "w", encoding="utf-8") as f:
        f.write(f"TOTAL TABLES: {len(doc.tables)}\n\n")
        for i, tbl in enumerate(doc.tables):
            f.write(f"=== TABLE {i} ===\n")
            f.write(f"Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)}\n")
            if len(tbl.rows) > 0:
                header = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
                f.write(f"Header: {' | '.join(header)}\n")
            if len(tbl.rows) > 1:
                sample = [c.text.strip().replace('\n', ' ') for c in tbl.rows[1].cells]
                f.write(f"Sample row 1: {' | '.join(sample)}\n")
            f.write("\n")

if __name__ == "__main__":
    dump_all_tables_info()

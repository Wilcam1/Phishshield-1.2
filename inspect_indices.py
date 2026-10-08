import docx

doc = docx.Document("Plantilla trabajo de grado 12 - v2.docx")

def print_p_range(start, end):
    for i in range(start, end):
        if i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            print(f"P{i:03d} | Style: '{p.style.name}' | Text: {repr(p.text[:60])}")

print("--- Paragraphs around Chapter 2 requirements and misplaced diagrams (385 to 455) ---")
print_p_range(385, 455)

print("\n--- Paragraphs around Chapter 3 and 4 (452 to 505) ---")
print_p_range(452, 505)

print("\n--- Paragraphs around Chapter 5 and Bibliografia (504 to 542) ---")
print_p_range(504, len(doc.paragraphs))

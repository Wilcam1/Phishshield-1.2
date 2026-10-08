import docx

def dump_document(filename, out_txt):
    doc = docx.Document(filename)
    with open(out_txt, "w", encoding="utf-8") as f:
        f.write(f"DOCUMENT: {filename}\n")
        f.write(f"SECTIONS: {len(doc.sections)}\n")
        f.write(f"TABLES: {len(doc.tables)}\n")
        f.write("="*80 + "\n\n")
        
        for i, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            style = p.style.name if p.style else ""
            if txt:
                f.write(f"[P{i:03d} | Style: {style}]\n{txt}\n\n")
                
        f.write("\n" + "="*80 + "\nTABLES DUMP:\n" + "="*80 + "\n")
        for t_idx, tbl in enumerate(doc.tables):
            f.write(f"\n--- TABLE {t_idx} (Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)}) ---\n")
            for r_idx, row in enumerate(tbl.rows):
                row_txt = [c.text.strip().replace("\n", " ") for c in row.cells]
                f.write(f"  Row {r_idx}: {' | '.join(row_txt)}\n")

if __name__ == "__main__":
    dump_document("Plantilla trabajo de grado 12.docx", "dump_plantilla_12.txt")
    dump_document("Proyecto_Fin_de_Grado_PhishShield_Compensar.docx", "dump_proyecto_draft.txt")
    print("Dumps created successfully.")

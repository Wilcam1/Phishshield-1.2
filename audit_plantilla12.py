import docx

def detailed_section_audit():
    doc = docx.Document("Plantilla trabajo de grado 12.docx")
    
    sections = [
        ("Portada", 0, 34),
        ("Dedicatoria y Agradecimientos", 35, 44),
        ("Resumen y Abstract", 45, 56),
        ("Contenido y Listas", 57, 129),
        ("Introducción", 130, 131),
        ("Antecedentes y Justificación", 132, 170),
        ("Objetivos", 171, 190),
        ("Alcances y Limitaciones", 191, 218),
        ("Cronograma y Presupuesto", 219, 224),
        ("Marco Teórico y Metodología", 225, 305),
        ("Estado del Arte y Marco Legal", 306, 326),
        ("Análisis del Proyecto", 327, 406),
        ("Modelado en Cap 2 (Error estructural)", 407, 451),
        ("Capítulo 3: Diseño del Proyecto", 452, 478),
        ("Capítulo 4: Implementación y Pruebas", 479, 503),
        ("Capítulo 5: Conclusiones y Recomendaciones", 504, 524),
        ("Bibliografía", 525, len(doc.paragraphs)-1)
    ]
    
    with open("section_audit_plantilla12.txt", "w", encoding="utf-8") as f:
        for name, start, end in sections:
            f.write(f"\n{'='*80}\n")
            f.write(f"SECCIÓN: {name} (Párrafos {start} a {end})\n")
            f.write(f"{'='*80}\n")
            for p_idx in range(start, min(end+1, len(doc.paragraphs))):
                p = doc.paragraphs[p_idx]
                t = p.text.strip()
                st = p.style.name if p.style else ""
                if t:
                    f.write(f"[P{p_idx:03d} | {st}]: {t}\n\n")

if __name__ == "__main__":
    detailed_section_audit()
    print("section_audit_plantilla12.txt written.")

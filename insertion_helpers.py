import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def apply_table_styling(table, headers, data, col_widths, alt_shading=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

    hdr = table.rows[0]
    hdr_trPr = hdr._element.get_or_add_trPr()
    hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
        tcPr.append(shd)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = 'Cambria'
            r.font.bold = True
            r.font.size = Pt(9.0)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row_values in enumerate(data, start=1):
        row = table.rows[r_idx]
        r_trPr = row._element.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg = "F8FAFC" if (r_idx % 2 == 0 and alt_shading) else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(val)
            tcPr = cell._element.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
            tcPr.append(shd)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            if c_idx == 0 and len(str(val)) <= 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Cambria'
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    for row in table.rows:
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = w

def add_p_after(ref_p, text, bold_prefix="", space_after=6, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, style_name="Normal"):
    doc = ref_p._parent
    new_p = doc.add_paragraph()
    new_p.style = style_name
    new_p.paragraph_format.line_spacing = 1.5
    new_p.paragraph_format.space_after = Pt(space_after)
    new_p.paragraph_format.alignment = align
    if bold_prefix:
        r_b = new_p.add_run(bold_prefix)
        r_b.font.name = 'Cambria'
        r_b.font.bold = True
        r_b.font.size = Pt(11)
    r_t = new_p.add_run(text)
    r_t.font.name = 'Cambria'
    r_t.font.size = Pt(11)
    r_t.font.italic = italic
    # Move new_p right after ref_p in XML
    ref_p._element.addnext(new_p._element)
    return new_p

def add_table_after(ref_p, headers, data, col_widths, alt_shading=True):
    doc = ref_p._parent
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    apply_table_styling(tbl, headers, data, col_widths, alt_shading)
    ref_p._element.addnext(tbl._element)
    return tbl

def add_figure_after(ref_p, img_path, caption_title, caption_note=""):
    if os.path.exists(img_path):
        doc = ref_p._parent
        
        # Paragraph for caption note
        p_note = None
        if caption_note:
            p_note = doc.add_paragraph()
            p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_note.paragraph_format.space_after = Pt(8)
            r_note = p_note.add_run(caption_note)
            r_note.font.name = 'Cambria'
            r_note.font.size = Pt(8.5)
            r_note.font.italic = True
            r_note.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Paragraph for caption title
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(caption_title)
        r_cap.font.name = 'Cambria'
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        # Paragraph for image
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(6.0))

        # Insert in order after ref_p:
        # ref_p -> p_img -> p_cap -> p_note
        if p_note:
            ref_p._element.addnext(p_note._element)
        ref_p._element.addnext(p_cap._element)
        ref_p._element.addnext(p_img._element)

print("Insertion helpers compiled.")

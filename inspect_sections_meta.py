import docx

doc = docx.Document("Plantilla trabajo de grado 12 - v2.docx")

print(f"Total sections: {len(doc.sections)}")
for i, s in enumerate(doc.sections):
    hdr = s.header
    ftr = s.footer
    hdr_txt = "".join(p.text for p in hdr.paragraphs) if hdr else ""
    ftr_txt = "".join(p.text for p in ftr.paragraphs) if ftr else ""
    print(f"Section {i:02d}: start_type={s.start_type}, header={repr(hdr_txt[:30])}, footer={repr(ftr_txt[:30])}")
